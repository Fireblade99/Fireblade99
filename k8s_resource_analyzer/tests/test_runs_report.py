"""End-to-end: fake VictoriaMetrics → analyzer → recommender → Excel/JSON."""

import json
import re
import time
from typing import Dict, List, Tuple

import pytest
from openpyxl import load_workbook

from src.config import Config
from src.core.analyzer import ResourceAnalyzer
from src.core.grouper import extract_base_name
from src.core.recommender import Recommender
from src.reporters import excel_reporter
from src.reporters.json_reporter import to_json

GI = 1024 ** 3
NS = "mpa-airflow-6679"
NOW = time.time()
H = 3600


class Pod:
    def __init__(self, name, start, hours, cpu, mem, req, lim=None, labels=None, oom=False):
        self.name, self.start, self.end = name, start, start + hours * H
        self.cpu, self.mem = cpu, mem          # (avg, peak) cores / bytes
        self.req, self.lim = req, lim or req   # (cpu cores, mem bytes)
        self.labels = labels or {}
        self.oom = oom

    def alive(self, a, b):
        return self.start <= b and self.end >= a


PODS = [
    # Two runs of the same Airflow task: asks 512Gi, uses ~6Gi
    Pod("gb-data-collect-task-ab12cd34", NOW - 3 * 86400, 2, (0.5, 0.8), (4 * GI, 6 * GI), (1.0, 512 * GI),
        labels={"label_dag_id": "gb_data", "label_task_id": "collect_task",
                "label_run_id": "scheduled__2026-10-06", "label_try_number": "1"}),
    Pod("gb-data-collect-task-zz98yy76", NOW - 1 * 86400, 3, (0.6, 0.9), (5 * GI, 7 * GI), (1.0, 512 * GI),
        labels={"label_dag_id": "gb_data", "label_task_id": "collect_task",
                "label_run_id": "scheduled__2026-10-08", "label_try_number": "1"}),
    # Under-provisioned and OOM-killed
    Pod("risky-task-qq11ww22", NOW - 2 * 86400, 1, (1.5, 2.0), (2 * GI, 3 * GI), (1.0, 2 * GI), oom=True),
    # Deployment replica, runs across chunk borders the whole week
    Pod("main-webserver-7d9f8b6c5d-abcde", NOW - 7 * 86400, 7 * 24, (0.01, 0.02), (1.5 * GI, 1.9 * GI),
        (0.5, 2 * GI)),
]


def _points(pod: Pod, attr: str, start: float, end: float, step: int) -> List[list]:
    avg, peak = getattr(pod, attr)
    out = []
    t = (int(start) // step) * step
    i = 0
    while t <= end:
        if pod.start <= t <= pod.end:
            # one peak sample, the rest chosen so the mean stays close to avg
            out.append([t, str(peak if i == 0 else avg)])
            i += 1
        t += step
    return out


class FakeResp:
    def __init__(self, result):
        self.status_code, self.ok, self.reason = 200, True, "OK"
        self._result = result
        self.text = ""

    def json(self):
        return {"status": "success", "data": {"result": self._result}}


class FakeVM:
    def __init__(self):
        self.calls: List[Tuple[str, Dict]] = []

    def post(self, url, data=None, timeout=None):
        endpoint = url.rsplit("/", 1)[-1]
        self.calls.append((endpoint, dict(data)))
        q = data["query"]
        if endpoint == "query_range":
            attr = "cpu" if "cpu_usage" in q else "mem"
            res = []
            for p in PODS:
                pts = _points(p, attr, data["start"], data["end"], data["step"])
                if pts:
                    res.append({"metric": {"namespace": NS, "pod": p.name, "container": "base"},
                                "values": pts})
            return FakeResp(res)

        # instant query evaluated at chunk end over the chunk window
        t = data["time"]
        win = int(re.search(r"\[(\d+)s\]", q).group(1))
        live = [p for p in PODS if p.alive(t - win, t)]
        res = []
        if "kube_pod_container_resource_" in q:
            idx = 0 if "requests" in q else 1
            for p in live:
                cpu, mem = (p.req, p.lim)[idx]
                for resource, val in (("cpu", cpu), ("memory", mem)):
                    res.append({"metric": {"namespace": NS, "pod": p.name, "container": "base",
                                           "resource": resource}, "value": [t, str(val)]})
        elif "kube_pod_labels" in q:
            for p in live:
                res.append({"metric": {"namespace": NS, "pod": p.name, **p.labels}, "value": [t, "1"]})
        elif "OOMKilled" in q:
            for p in live:
                if p.oom:
                    res.append({"metric": {"namespace": NS, "pod": p.name, "container": "base"},
                                "value": [t, "1"]})
        return FakeResp(res)


@pytest.fixture
def analysed(monkeypatch):
    vm = FakeVM()
    monkeypatch.setattr("requests.Session.post", lambda self, *a, **kw: vm.post(*a, **kw))
    cfg = Config.default()
    cfg.kubernetes.enabled = False
    cfg.prometheus.url = "http://vm/select/0/prometheus"
    cfg.prometheus.pause_seconds = 0
    cfg.analysis.namespaces = [NS]
    groups = ResourceAnalyzer(cfg).analyze()
    recs = Recommender().process_all(groups)
    return vm, groups, recs


def test_queries_are_chunked_and_aggregated(analysed):
    vm, _, _ = analysed
    ranges = [d for e, d in vm.calls if e == "query_range"]
    # 7 days in 24h chunks → 7 (or 8 with a partial edge) requests per usage metric
    assert 14 <= len(ranges) <= 16
    for d in ranges:
        assert d["end"] - d["start"] <= 24 * H
        assert d["query"].startswith("max by (namespace, pod, container)")
        assert f'namespace="{NS}"' in d["query"]
    assert all("[$window]" not in d["query"] for _, d in vm.calls)


def test_runs_and_groups(analysed):
    _, groups, recs = analysed
    by_name = {g.base_name: g for g in groups}
    assert set(by_name) == {"gb-data-collect-task", "risky-task", "main-webserver"}

    gb = by_name["gb-data-collect-task"]
    assert len(gb.runs) == 2
    assert gb.dag_id == "gb_data" and gb.task_id == "collect_task"
    assert [r.run_id for r in gb.runs] == ["scheduled__2026-10-06", "scheduled__2026-10-08"]
    assert gb.memory_request == 512 * GI
    assert gb.max_memory_usage == 7 * GI
    run = gb.runs[0]
    assert run.mem_over == 512 * GI - 6 * GI
    assert run.mem_over_ratio == pytest.approx(1 - 6 / 512)
    assert 2 <= run.duration_hours <= 2.2
    # ~ (512 − ~4) GiB × ~2h idle
    assert 900 < run.mem_idle_gib_hours < 1100

    web = by_name["main-webserver"]
    assert len(web.runs) == 1  # one pod across all 7 chunks
    assert web.runs[0].duration_hours > 7 * 24 - 1

    risky = next(r for r in recs if r.group.base_name == "risky-task")
    assert risky.is_risky
    assert any("OOMKilled in 1 of 1 runs" in x for x in risky.reasons)
    wasteful = next(r for r in recs if r.group.base_name == "gb-data-collect-task")
    assert wasteful.is_wasteful


def test_excel_and_json(analysed, tmp_path):
    _, _, recs = analysed
    path = tmp_path / "report.xlsx"
    excel_reporter.generate(recs, output_path=str(path), show_only_waste=False, lookback_days=7)

    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Recommendations", "Runs"]

    runs = wb["Runs"]
    header = [c.value for c in runs[1]]
    rows = [dict(zip(header, [c.value for c in r])) for r in runs.iter_rows(min_row=2)]
    assert len(rows) == 4
    gb = [r for r in rows if r["Workload"] == "gb-data-collect-task"]
    assert {r["Run ID"] for r in gb} == {"scheduled__2026-10-06", "scheduled__2026-10-08"}
    assert all(r["Mem Request, GiB"] == 512 for r in gb)
    assert max(r["Mem Max, GiB"] for r in gb) == 7
    assert any(r["OOMKilled"] == "YES" for r in rows)

    summary = [[c.value for c in r] for r in wb["Summary"].iter_rows()]
    labels = [r[0] for r in summary]
    assert "Memory reserved but idle" in labels
    assert any(str(x).startswith("Top 10 — Idle Memory Reservation") for x in labels)

    data = json.loads(to_json(recs, show_only_waste=False))
    gbj = next(x for x in data if x["workload"] == "gb-data-collect-task")
    assert gbj["runs"] == 2 and gbj["dag_id"] == "gb_data"
    assert gbj["resource_hours"]["memory_idle_gib_hours"] > 2000


def test_proxy_env_ignored_by_default():
    from src.clients.prom_client import PrometheusClient

    assert PrometheusClient("http://x").__dict__["_session"].trust_env is False
    assert PrometheusClient("http://x", use_proxy=True).__dict__["_session"].trust_env is True


@pytest.mark.parametrize("pod,base", [
    ("gb-data-collect-task-ab12cd34", "gb-data-collect-task"),
    ("main-webserver-7d9f8b6c5d-abcde", "main-webserver"),
    ("redis-0", "redis"),
])
def test_base_name(pod, base):
    assert extract_base_name(pod) == base


def _group(cpu_req, cpu_max, mem_req, mem_max, oom=False):
    from src.core.grouper import WorkloadGroup
    from src.core.runs import PodRun

    g = WorkloadGroup(namespace=NS, base_name="w", container="base")
    g.cpu_request, g.max_cpu_usage, g.memory_request, g.max_memory_usage = cpu_req, cpu_max, mem_req, mem_max
    g.runs = [PodRun(NS, "w-abcdefgh", "base", "w", 0, 3600, oom_killed=oom)]
    return g


@pytest.mark.parametrize("cpu_req,cpu_max,mem_req,mem_max,oom,cpu_act,mem_act,risky", [
    # CPU 4% above request is within the 20% tolerance; memory 40% unused is below the 50% threshold
    (1.0, 1.04, 60 * GI, 30 * GI * 1.2, False, "ok", "ok", False),
    (1.0, 1.30, 60 * GI, 2 * GI, False, "up", "down", True),
    (4.0, 0.5, 8 * GI, 9 * GI, False, "down", "up", True),
    (1.0, 0.9, 8 * GI, 4 * GI, True, "ok", "up", True),
])
def test_actions_per_resource(cpu_req, cpu_max, mem_req, mem_max, oom, cpu_act, mem_act, risky):
    rec = Recommender().recommend(_group(cpu_req, cpu_max, mem_req, mem_max, oom))
    assert (rec.cpu_action, rec.memory_action, rec.is_risky) == (cpu_act, mem_act, risky)

