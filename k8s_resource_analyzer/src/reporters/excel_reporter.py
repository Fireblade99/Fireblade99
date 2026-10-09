"""
Excel reporter – generates a formatted .xlsx report for business stakeholders.

Sheets
------
1. Summary          – total waste stats, requested vs used resource-hours,
                      top-10 wasteful workloads, top-10 idle reservations
2. Recommendations  – one row per workload with colour-coded waste %
3. Runs             – one row per run (pod) with request vs actual usage
"""

import datetime
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, GradientFill, PatternFill, Side
)
from openpyxl.utils import get_column_letter

from ..core.recommender import Recommendation
from ..core.runs import GIB
from ..utils import fmt_bytes, fmt_cores

# ── Colour palette ─────────────────────────────────────────────────────────
_RED    = "FFDD2222"
_ORANGE = "FFFF6600"
_YELLOW = "FFFFC000"
_GREEN  = "FF00AA44"
_LIGHT_RED    = "FFFFC7C7"
_LIGHT_ORANGE = "FFFFDDB0"
_LIGHT_YELLOW = "FFFFF0B0"
_LIGHT_GREEN  = "FFD6F5D6"
_HEADER_BG    = "FF1F4E79"
_SUMMARY_BG   = "FF2E75B6"
_RISKY_BG     = "FFFFF2CC"


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(bold: bool = False, color: str = "FF000000", size: int = 11) -> Font:
    return Font(bold=bold, color=color, size=size)


def _border() -> Border:
    thin = Side(style="thin", color="FFD0D0D0")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _waste_fill(ratio: Optional[float]) -> Optional[PatternFill]:
    if ratio is None or ratio <= 0:
        return None
    if ratio >= 0.80:
        return _fill(_LIGHT_RED)
    if ratio >= 0.60:
        return _fill(_LIGHT_ORANGE)
    if ratio >= 0.40:
        return _fill(_LIGHT_YELLOW)
    return _fill(_LIGHT_GREEN)


def _waste_label(ratio: Optional[float]) -> str:
    if ratio is None or ratio <= 0:
        return ""
    if ratio >= 0.80:
        return f"{ratio:.0%} ⚠ HIGH"
    if ratio >= 0.60:
        return f"{ratio:.0%} ↑ MED"
    return f"{ratio:.0%}"


# ── Public API ─────────────────────────────────────────────────────────────

def generate(
    recommendations: List[Recommendation],
    output_path: str,
    show_only_waste: bool = True,
    lookback_days: int = 7,
) -> str:
    """Write the report to *output_path* and return the path."""
    all_items = list(recommendations)
    items = all_items
    if show_only_waste:
        items = [r for r in items if r.is_wasteful or r.is_risky]
    items.sort(
        key=lambda r: (r.memory_waste_bytes or 0) if (r.memory_waste_bytes or 0) > 0 else 0,
        reverse=True,
    )

    wb = Workbook()
    _sheet_summary(wb, items, lookback_days, all_items)
    _sheet_report(wb, items)
    _sheet_runs(wb, items)
    wb.save(output_path)
    return output_path


# ── Sheet 1: Summary ───────────────────────────────────────────────────────

def _sheet_summary(
    wb: Workbook,
    items: List[Recommendation],
    lookback_days: int,
    all_items: Optional[List[Recommendation]] = None,
) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 22

    date_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    wasteful = [r for r in items if r.is_wasteful]
    risky    = [r for r in items if r.is_risky]
    total_mem = sum((r.memory_waste_bytes or 0) for r in wasteful if (r.memory_waste_bytes or 0) > 0)
    total_cpu = sum((r.cpu_waste_cores   or 0) for r in wasteful if (r.cpu_waste_cores   or 0) > 0)

    # Resource-hours over every analysed run (not only the flagged ones)
    runs = [run for r in (all_items if all_items is not None else items) for run in r.group.runs]
    cpu_req_h = sum(x.cpu_requested_core_hours for x in runs)
    cpu_idle_h = sum(x.cpu_idle_core_hours for x in runs)
    mem_req_h = sum(x.mem_requested_gib_hours for x in runs)
    mem_idle_h = sum(x.mem_idle_gib_hours for x in runs)

    def _hdr(row: int, text: str) -> None:
        c = ws.cell(row=row, column=1, value=text)
        c.font = _font(bold=True, color="FFFFFFFF", size=13)
        c.fill = _fill(_SUMMARY_BG)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[row].height = 22
        ws.merge_cells(f"A{row}:B{row}")

    def _row(row: int, label: str, value: str, bold_val: bool = False) -> None:
        lc = ws.cell(row=row, column=1, value=label)
        vc = ws.cell(row=row, column=2, value=value)
        lc.font = _font(size=11)
        vc.font = _font(bold=bold_val, size=11)
        lc.border = vc.border = _border()
        vc.alignment = Alignment(horizontal="right")
        ws.row_dimensions[row].height = 18

    # Title
    title = ws.cell(row=1, column=1, value="☁  Kubernetes Resource Waste Report")
    title.font = _font(bold=True, color="FFFFFFFF", size=16)
    title.fill = _fill(_HEADER_BG)
    title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A1:B1")

    ws.cell(row=2, column=1, value=f"Generated: {date_str}  |  Lookback: {lookback_days} days")
    ws.cell(row=2, column=1).font = _font(color="FF666666", size=10)
    ws.merge_cells("A2:B2")
    ws.row_dimensions[2].height = 16

    ws.row_dimensions[3].height = 8  # spacer

    _hdr(4, "Overview")
    _row(5,  "Analysis period",          f"{lookback_days} days")
    _row(6,  "Over-provisioned workloads", str(len(wasteful)), bold_val=True)
    _row(7,  "Under-provisioned (risky)",  str(len(risky)),    bold_val=True)
    _row(8,  "Total wasted memory",        fmt_bytes(total_mem), bold_val=True)
    _row(9,  "Total wasted CPU",           f"{fmt_cores(total_cpu)} cores", bold_val=True)
    _row(10, "Runs (pods) analysed",       str(len(runs)))
    _row(11, "CPU requested × time",       f"{cpu_req_h:,.0f} core·h")
    _row(12, "CPU reserved but idle",      f"{cpu_idle_h:,.0f} core·h ({_share(cpu_idle_h, cpu_req_h)})",
         bold_val=True)
    _row(13, "Memory requested × time",    f"{mem_req_h:,.0f} GiB·h")
    _row(14, "Memory reserved but idle",   f"{mem_idle_h:,.0f} GiB·h ({_share(mem_idle_h, mem_req_h)})",
         bold_val=True)

    ws.row_dimensions[15].height = 8  # spacer

    _hdr(16, "Top 10 — Highest Memory Waste")
    top10_headers = ["Workload", "Namespace", "Container",
                     "Mem Request", "Mem Max", "Mem Waste", "Waste %", "Recommended Request"]
    col_widths = [42, 24, 20, 14, 14, 14, 10, 22]
    for i, h in enumerate(top10_headers, start=1):
        c = ws.cell(row=17, column=i, value=h)
        c.font = _font(bold=True, color="FFFFFFFF", size=10)
        c.fill = _fill(_SUMMARY_BG)
        c.border = _border()
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = col_widths[i - 1]

    for idx, rec in enumerate(items[:10], start=18):
        g = rec.group
        row_data = [
            g.base_name,
            g.namespace,
            g.container,
            fmt_bytes(g.memory_request),
            fmt_bytes(g.max_memory_usage),
            fmt_bytes(rec.memory_waste_bytes) if (rec.memory_waste_bytes or 0) > 0 else "-",
            _waste_label(rec.memory_waste_ratio),
            fmt_bytes(rec.recommended_memory_request),
        ]
        fill = _waste_fill(rec.memory_waste_ratio) or PatternFill()
        for col, val in enumerate(row_data, start=1):
            c = ws.cell(row=idx, column=col, value=val)
            c.font = _font(size=10)
            c.border = _border()
            c.alignment = Alignment(horizontal="left" if col <= 3 else "right")
            if col in (6, 7):
                c.fill = fill

    # Top 10 by reserved-but-idle memory over the window (frequency × duration × over-request)
    start = 18 + min(len(items), 10) + 1
    _hdr(start, f"Top 10 — Idle Memory Reservation over {lookback_days} days")
    idle_headers = ["Workload", "DAG", "Runs", "Runtime, h",
                    "Mem Request", "Mem Max", "Idle Mem GiB·h", "Idle CPU core·h"]
    for i, h in enumerate(idle_headers, start=1):
        c = ws.cell(row=start + 1, column=i, value=h)
        c.font = _font(bold=True, color="FFFFFFFF", size=10)
        c.fill = _fill(_SUMMARY_BG)
        c.border = _border()
        c.alignment = Alignment(horizontal="center")
    by_idle = sorted(items, key=lambda r: r.group.mem_idle_gib_hours, reverse=True)[:10]
    for idx, rec in enumerate(by_idle, start=start + 2):
        g = rec.group
        row_data = [
            g.base_name,
            g.dag_id,
            len(g.runs),
            round(g.runtime_hours, 1),
            fmt_bytes(g.memory_request),
            fmt_bytes(g.max_memory_usage),
            round(g.mem_idle_gib_hours, 1),
            round(g.cpu_idle_core_hours, 1),
        ]
        for col, val in enumerate(row_data, start=1):
            c = ws.cell(row=idx, column=col, value=val)
            c.font = _font(size=10)
            c.border = _border()
            c.alignment = Alignment(horizontal="left" if col <= 2 else "right")
            if col == 7:
                c.fill = _fill(_LIGHT_RED)


def _share(part: float, whole: float) -> str:
    return f"{part / whole:.0%}" if whole else "-"


# ── Sheet 2: Full report ───────────────────────────────────────────────────

_REPORT_HEADERS = [
    "Namespace", "Workload", "DAG", "Task", "Container", "Runs",
    "CPU Request", "CPU Avg", "CPU Max", "CPU Rec",
    "Mem Request", "Mem Avg", "Mem Max", "Mem Rec",
    "CPU Waste", "CPU Waste %",
    "Mem Waste", "Mem Waste %",
    "Idle CPU core·h", "Idle Mem GiB·h", "OOM runs",
    "Status", "Notes",
]
_REPORT_WIDTHS = [24, 48, 28, 28, 14, 7, 12, 12, 12, 12, 14, 14, 14, 14, 12, 12, 14, 12, 12, 13, 9, 14, 50]
# 1-based columns that get special treatment
_C_CPU_WASTE = (15, 16)
_C_MEM_WASTE = (17, 18)
_C_IDLE_MEM = 20
_C_OOM = 21
_C_LEFT = (1, 2, 3, 4, 5, 22, 23)
_C_NOTES = 23


def _sheet_report(wb: Workbook, items: List[Recommendation]) -> None:
    ws = wb.create_sheet(title="Recommendations")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_REPORT_HEADERS))}{max(1, len(items) + 1)}"

    for i, (h, w) in enumerate(zip(_REPORT_HEADERS, _REPORT_WIDTHS), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = _font(bold=True, color="FFFFFFFF", size=10)
        c.fill = _fill(_HEADER_BG)
        c.border = _border()
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 28

    for row_idx, rec in enumerate(items, start=2):
        g = rec.group
        mem_waste_val = rec.memory_waste_bytes if (rec.memory_waste_bytes or 0) > 0 else None
        cpu_waste_val = rec.cpu_waste_cores   if (rec.cpu_waste_cores   or 0) > 0 else None

        if rec.is_risky and rec.is_wasteful:
            status = "⚠ RISKY + WASTE"
        elif rec.is_risky:
            status = "⚠ UNDER-PROVISIONED"
        elif rec.is_wasteful:
            status = "↓ OVER-PROVISIONED"
        else:
            status = "✓ OK"

        row_data = [
            g.namespace,
            g.base_name,
            g.dag_id,
            g.task_id,
            g.container,
            len(g.runs) or len(g.pod_names),
            fmt_cores(g.cpu_request),
            fmt_cores(g.avg_cpu_usage),
            fmt_cores(g.max_cpu_usage),
            fmt_cores(rec.recommended_cpu_request),
            fmt_bytes(g.memory_request),
            fmt_bytes(g.avg_memory_usage),
            fmt_bytes(g.max_memory_usage),
            fmt_bytes(rec.recommended_memory_request),
            fmt_cores(cpu_waste_val),
            f"{rec.cpu_waste_ratio:.0%}" if rec.cpu_waste_ratio and rec.cpu_waste_ratio > 0 else "-",
            fmt_bytes(mem_waste_val),
            _waste_label(rec.memory_waste_ratio) if (rec.memory_waste_bytes or 0) > 0 else "-",
            round(g.cpu_idle_core_hours, 1),
            round(g.mem_idle_gib_hours, 1),
            g.oom_runs or "",
            status,
            " | ".join(rec.reasons),
        ]

        mem_fill  = _waste_fill(rec.memory_waste_ratio)
        cpu_fill  = _waste_fill(rec.cpu_waste_ratio)
        risky_row = _fill(_RISKY_BG) if rec.is_risky and not rec.is_wasteful else None

        for col, val in enumerate(row_data, start=1):
            c = ws.cell(row=row_idx, column=col, value=val)
            c.font = _font(size=10)
            c.border = _border()
            c.alignment = Alignment(
                horizontal="left" if col in _C_LEFT else "right",
                wrap_text=(col == _C_NOTES),
            )
            # Apply colour
            if col in _C_CPU_WASTE and cpu_fill:
                c.fill = cpu_fill
            elif col in _C_MEM_WASTE and mem_fill:
                c.fill = mem_fill
            elif col == _C_OOM and g.oom_runs:
                c.fill = _fill(_LIGHT_RED)
            elif risky_row and col not in _C_CPU_WASTE + _C_MEM_WASTE:
                c.fill = risky_row

        ws.row_dimensions[row_idx].height = 16


# ── Sheet 3: Runs ──────────────────────────────────────────────────────────

# Numbers stay numeric (cores / GiB / hours) so the sheet can be sorted and filtered
_RUN_COLUMNS = [
    # header, width, number format
    ("Namespace", 22, None),
    ("Workload", 44, None),
    ("DAG", 28, None),
    ("Task", 28, None),
    ("Run ID", 30, None),
    ("Try", 5, None),
    ("Pod", 50, None),
    ("Container", 12, None),
    ("Start", 17, "yyyy-mm-dd hh:mm"),
    ("End", 17, "yyyy-mm-dd hh:mm"),
    ("Duration, h", 10, "0.00"),
    ("CPU Request", 10, "0.000"),
    ("CPU Limit", 10, "0.000"),
    ("CPU Avg", 10, "0.000"),
    ("CPU Max", 10, "0.000"),
    ("CPU Over-request", 11, "0.000"),
    ("CPU Over %", 9, "0%"),
    ("Mem Request, GiB", 11, "0.00"),
    ("Mem Limit, GiB", 11, "0.00"),
    ("Mem Avg, GiB", 11, "0.00"),
    ("Mem Max, GiB", 11, "0.00"),
    ("Mem Over-request, GiB", 12, "0.00"),
    ("Mem Over %", 9, "0%"),
    ("Idle CPU core·h", 10, "0.00"),
    ("Idle Mem GiB·h", 10, "0.00"),
    ("OOMKilled", 10, None),
]
_RUN_C_CPU_OVER = (16, 17)
_RUN_C_MEM_OVER = (22, 23)


def _gib(b: Optional[float]) -> Optional[float]:
    return None if b is None else b / GIB


def _sheet_runs(wb: Workbook, items: List[Recommendation]) -> None:
    ws = wb.create_sheet(title="Runs")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "C2"

    for i, (h, w, _) in enumerate(_RUN_COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = _font(bold=True, color="FFFFFFFF", size=10)
        c.fill = _fill(_HEADER_BG)
        c.border = _border()
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30

    row_idx = 1
    for rec in items:
        for run in sorted(rec.group.runs, key=lambda x: x.start, reverse=True):
            row_idx += 1
            row_data = [
                run.namespace,
                run.base_name,
                run.dag_id,
                run.task_id,
                run.run_id,
                run.try_number,
                run.pod,
                run.container,
                run.started_at,
                run.ended_at,
                run.duration_hours,
                run.cpu_request,
                run.cpu_limit,
                run.cpu_avg,
                run.cpu_max,
                run.cpu_over,
                run.cpu_over_ratio,
                _gib(run.mem_request),
                _gib(run.mem_limit),
                _gib(run.mem_avg),
                _gib(run.mem_max),
                _gib(run.mem_over),
                run.mem_over_ratio,
                run.cpu_idle_core_hours,
                run.mem_idle_gib_hours,
                "YES" if run.oom_killed else "",
            ]
            cpu_fill = _waste_fill(run.cpu_over_ratio)
            mem_fill = _waste_fill(run.mem_over_ratio)
            for col, val in enumerate(row_data, start=1):
                c = ws.cell(row=row_idx, column=col, value=val)
                c.font = _font(size=10)
                c.border = _border()
                fmt = _RUN_COLUMNS[col - 1][2]
                if fmt:
                    c.number_format = fmt
                if col in _RUN_C_CPU_OVER and cpu_fill:
                    c.fill = cpu_fill
                elif col in _RUN_C_MEM_OVER and mem_fill:
                    c.fill = mem_fill
                elif col == len(_RUN_COLUMNS) and run.oom_killed:
                    c.fill = _fill(_LIGHT_RED)
                # Under-request (used more than asked) → risky colour
                if col in (_RUN_C_CPU_OVER[0], _RUN_C_MEM_OVER[0]) and val is not None and val < 0:
                    c.fill = _fill(_RISKY_BG)

    ws.auto_filter.ref = f"A1:{get_column_letter(len(_RUN_COLUMNS))}{max(1, row_idx)}"
